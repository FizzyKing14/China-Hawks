#!/usr/bin/env python3
"""Generate a car-flipping profit tracker as a real .xlsx with live formulas.

Run:  python3 build_excel.py
Out:  车辆利润表.xlsx
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = "车辆利润表.xlsx"
DATA_ROWS = 60          # how many blank car rows to pre-build with formulas

# ---- styling helpers -------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")   # yellow = type here
CALC_FILL = PatternFill("solid", fgColor="E2EFDA")    # green = auto-calculated
TITLE_FONT = Font(bold=True, size=16, color="1F4E78")
LABEL_FONT = Font(bold=True, size=11)
MONEY = '#,##0.00'
PCT = '0.0%'
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


# ---- columns: (header, key, kind, width) -----------------------------------
# kind: "input" = you type it, "calc" = a formula fills it.
COLUMNS = [
    ("编号", "id", "input", 8),
    ("车型 / 描述", "desc", "input", 22),
    ("进车价", "buy", "input", 12),
    ("修车/翻新", "recon", "input", 12),
    ("运费", "ship", "input", 10),
    ("过户/上牌/检测", "fees", "input", 14),
    ("其他杂费", "misc", "input", 11),
    ("总成本", "cost", "calc", 12),
    ("卖出价", "sell", "input", 12),
    ("利润", "profit", "calc", 12),
    ("利润率", "margin", "calc", 10),
    ("回报率 ROI", "roi", "calc", 11),
    ("状态", "status", "calc", 10),
]
COL = {c[1]: i + 1 for i, c in enumerate(COLUMNS)}   # key -> 1-based col index


def cl(key: str) -> str:
    """Column letter for a data key."""
    return get_column_letter(COL[key])


def build() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "车辆利润"

    # Title row
    ws.merge_cells("A1:M1")
    ws["A1"] = "🚗 车辆买卖利润表"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells("A2:M2")
    ws["A2"] = ("黄色格子=自己填  |  绿色格子=自动算  |  "
                "卖出价留空表示还没卖出（在售）")
    ws["A2"].font = Font(italic=True, size=10, color="808080")

    header_row = 4
    first_data = header_row + 1
    last_data = first_data + DATA_ROWS - 1

    # Header
    for c in COLUMNS:
        col_idx = COL[c[1]]
        cell = ws.cell(row=header_row, column=col_idx, value=c[0])
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = c[3]
    ws.row_dimensions[header_row].height = 30

    # Data rows with formulas
    for r in range(first_data, last_data + 1):
        for c in COLUMNS:
            key = c[1]
            cell = ws.cell(row=r, column=COL[key])
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")

            if key == "cost":
                cell.value = (f"=IF(COUNT({cl('buy')}{r}:{cl('misc')}{r})=0,\"\","
                              f"SUM({cl('buy')}{r}:{cl('misc')}{r}))")
                cell.number_format = MONEY
                cell.fill = CALC_FILL
            elif key == "profit":
                cell.value = (f"=IF(OR({cl('sell')}{r}=\"\",{cl('cost')}{r}=\"\"),\"\","
                              f"{cl('sell')}{r}-{cl('cost')}{r})")
                cell.number_format = MONEY
                cell.fill = CALC_FILL
            elif key == "margin":
                # profit / sell  (利润占卖价比例)
                cell.value = (f"=IF(OR({cl('sell')}{r}=\"\",{cl('sell')}{r}=0),\"\","
                              f"{cl('profit')}{r}/{cl('sell')}{r})")
                cell.number_format = PCT
                cell.fill = CALC_FILL
            elif key == "roi":
                # profit / cost  (利润占投入比例)
                cell.value = (f"=IF(OR({cl('cost')}{r}=\"\",{cl('cost')}{r}=0),\"\","
                              f"{cl('profit')}{r}/{cl('cost')}{r})")
                cell.number_format = PCT
                cell.fill = CALC_FILL
            elif key == "status":
                cell.value = (f"=IF({cl('cost')}{r}=\"\",\"\","
                              f"IF({cl('sell')}{r}=\"\",\"在售\",\"已售\"))")
                cell.fill = CALC_FILL
            else:
                # input cells
                cell.fill = INPUT_FILL
                if key in ("buy", "recon", "ship", "fees", "misc", "sell"):
                    cell.number_format = MONEY

    # ---- summary block ----------------------------------------------------
    s = last_data + 3
    ws.cell(row=s, column=1, value="📊 汇总").font = Font(bold=True, size=14,
                                                         color="1F4E78")

    def summary(label, formula, fmt=MONEY, offset=1):
        row = s + offset
        lc = ws.cell(row=row, column=1, value=label)
        lc.font = LABEL_FONT
        vc = ws.cell(row=row, column=3, value=formula)
        vc.number_format = fmt
        vc.font = Font(bold=True)
        vc.fill = CALC_FILL
        vc.border = BORDER
        return row

    prof = f"{cl('profit')}{first_data}:{cl('profit')}{last_data}"
    cost = f"{cl('cost')}{first_data}:{cl('cost')}{last_data}"
    stat = f"{cl('status')}{first_data}:{cl('status')}{last_data}"
    sell = f"{cl('sell')}{first_data}:{cl('sell')}{last_data}"

    summary("总台数（已录入）", f'=COUNTIF({stat},"<>")', fmt='0', offset=1)
    summary("已售台数", f'=COUNTIF({stat},"已售")', fmt='0', offset=2)
    summary("在售台数", f'=COUNTIF({stat},"在售")', fmt='0', offset=3)
    summary("已售总利润", f'=SUM({prof})', offset=4)
    summary("已售平均每台利润",
            f'=IF(COUNTIF({stat},"已售")=0,0,SUM({prof})/COUNTIF({stat},"已售"))',
            offset=5)
    # ROI overall = total profit / total cost of SOLD cars
    sold_cost = (f'=SUMIFS({cost},{stat},"已售")')
    summary("已售总成本", sold_cost, offset=6)
    summary("整体回报率 ROI",
            f'=IF(SUMIFS({cost},{stat},"已售")=0,0,'
            f'SUM({prof})/SUMIFS({cost},{stat},"已售"))',
            fmt=PCT, offset=7)
    summary("在售占用资金（成本）", f'=SUMIFS({cost},{stat},"在售")', offset=8)

    # Freeze header so it stays visible while scrolling
    ws.freeze_panes = f"A{first_data}"

    # An example row so it's obvious how to use
    ex = first_data
    ws.cell(row=ex, column=COL["id"], value="示例")
    ws.cell(row=ex, column=COL["desc"], value="2018 起亚 K5")
    ws.cell(row=ex, column=COL["buy"], value=8000)
    ws.cell(row=ex, column=COL["recon"], value=1200)
    ws.cell(row=ex, column=COL["ship"], value=300)
    ws.cell(row=ex, column=COL["fees"], value=400)
    ws.cell(row=ex, column=COL["misc"], value=200)
    ws.cell(row=ex, column=COL["sell"], value=12500)

    wb.save(OUT)
    print(f"已生成 {OUT}")


if __name__ == "__main__":
    build()
